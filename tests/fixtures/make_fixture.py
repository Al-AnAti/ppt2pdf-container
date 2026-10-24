"""Build tests/fixtures/animated.pptx: a deck with animations and hidden shapes.

Run once (needs python-pptx + lxml):  python tests/fixtures/make_fixture.py

Slide 1 exercises every case flatten.py handles:
  ALWAYS VISIBLE       plain shape, no animation
  HIDDEN SHAPE         hidden="1" in the XML, no animation  -> must be revealed
  ENTRANCE STAYS       main-sequence entrance effect         -> must stay
  EXIT GONE            entrance then exit in main sequence   -> must be removed
  CLICK TRIGGERED      entrance in an interactive sequence   -> must stay
  PARAGRAPH EXIT       paragraph-level exit effect           -> must stay (whole shape kept)
Slide 2 has a transition and an emphasis effect only (EMPHASIS STAYS).
"""
from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE / "animated.pptx"
P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"
NSMAP = {"p": P_NS}

TEXTS = [
    "ALWAYS VISIBLE",
    "HIDDEN SHAPE",
    "ENTRANCE STAYS",
    "EXIT GONE",
    "CLICK TRIGGERED",
    "PARAGRAPH EXIT",
]


def build_deck(path: Path) -> None:
    prs = Presentation()
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    for i, text in enumerate(TEXTS):
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4 + i * 0.9), Inches(6), Inches(0.7))
        box.name = text
        para = box.text_frame.paragraphs[0]
        para.text = text
        para.runs[0].font.size = Pt(28)
        if text == "PARAGRAPH EXIT":
            box.text_frame.add_paragraph().text = "second paragraph"

    slide2 = prs.slides.add_slide(blank)
    box = slide2.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(6), Inches(0.7))
    box.name = "EMPHASIS STAYS"
    box.text_frame.text = "EMPHASIS STAYS"

    prs.save(path)


_ids = iter(range(100, 10_000))


def _effect(preset_class: str, spid: str, paragraph: bool = False, node_type: str = "clickEffect") -> str:
    to_val = "hidden" if preset_class == "exit" else "visible"
    tgt = f'<p:spTgt spid="{spid}"><p:txEl><p:pRg st="0" end="0"/></p:txEl></p:spTgt>' if paragraph else f'<p:spTgt spid="{spid}"/>'
    if preset_class == "emph":
        body = (
            f'<p:animClr clrSpc="rgb" dir="cw"><p:cBhvr><p:cTn id="{next(_ids)}" dur="500" fill="hold"/>'
            f"<p:tgtEl>{tgt}</p:tgtEl><p:attrNameLst><p:attrName>style.color</p:attrName></p:attrNameLst></p:cBhvr>"
            '<p:to><p:srgbClr val="FF0000"/></p:to></p:animClr>'
        )
    else:
        body = (
            f'<p:set><p:cBhvr><p:cTn id="{next(_ids)}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>'
            f"<p:tgtEl>{tgt}</p:tgtEl><p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>"
            f'<p:to><p:strVal val="{to_val}"/></p:to></p:set>'
        )
    return (
        f'<p:par><p:cTn id="{next(_ids)}" presetID="1" presetClass="{preset_class}" presetSubtype="0" fill="hold" grpId="0" nodeType="{node_type}">'
        f'<p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{body}</p:childTnLst></p:cTn></p:par>'
    )


def _click_group(*effects: str) -> str:
    return (
        f'<p:par><p:cTn id="{next(_ids)}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/></p:stCondLst><p:childTnLst>'
        f'<p:par><p:cTn id="{next(_ids)}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
        + "".join(effects)
        + "</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>"
    )


def _timing(main_groups: list[str], interactive: str = "") -> str:
    return (
        f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
        f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
        + "".join(main_groups)
        + '</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
        '<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
        + interactive
        + "</p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>"
    )


def _interactive(trigger_spid: str, *effects: str) -> str:
    cond = f'<p:cond evt="onClick" delay="0"><p:tgtEl><p:spTgt spid="{trigger_spid}"/></p:tgtEl></p:cond>'
    return (
        f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="{next(_ids)}" restart="whenNotActive" fill="hold" evtFilter="cancelBubble" nodeType="interactiveSeq">'
        f"<p:stCondLst>{cond}</p:stCondLst>"
        '<p:endSync evt="end" delay="0"><p:rtn val="all"/></p:endSync>'
        f'<p:childTnLst>{_click_group(*effects)}</p:childTnLst></p:cTn><p:nextCondLst>{cond}</p:nextCondLst></p:seq>'
    )


def inject(path: Path) -> None:
    """Add timing trees, a transition and a hidden flag to the python-pptx deck."""
    with zipfile.ZipFile(path) as zin:
        parts = {info.filename: zin.read(info) for info in zin.infolist()}
        infos = zin.infolist()

    def ids(xml: bytes) -> dict[str, str]:
        root = etree.fromstring(xml)
        return {el.get("name"): el.get("id") for el in root.iter(f"{{{P_NS}}}cNvPr")}

    s1 = etree.fromstring(parts["ppt/slides/slide1.xml"])
    id1 = ids(parts["ppt/slides/slide1.xml"])
    for cNvPr in s1.iter(f"{{{P_NS}}}cNvPr"):
        if cNvPr.get("name") == "HIDDEN SHAPE":
            cNvPr.set("hidden", "1")
    timing1 = _timing(
        [
            _click_group(_effect("entr", id1["ENTRANCE STAYS"]), _effect("entr", id1["EXIT GONE"], node_type="withEffect")),
            _click_group(_effect("entr", id1["PARAGRAPH EXIT"])),
            _click_group(_effect("exit", id1["EXIT GONE"]), _effect("exit", id1["PARAGRAPH EXIT"], paragraph=True, node_type="withEffect")),
        ],
        _interactive(id1["ALWAYS VISIBLE"], _effect("entr", id1["CLICK TRIGGERED"])),
    )
    s1.append(etree.fromstring(timing1))
    parts["ppt/slides/slide1.xml"] = etree.tostring(s1, xml_declaration=True, encoding="UTF-8", standalone=True)

    s2 = etree.fromstring(parts["ppt/slides/slide2.xml"])
    id2 = ids(parts["ppt/slides/slide2.xml"])
    cSld = s2.find(f"{{{P_NS}}}cSld")
    cSld.addnext(etree.fromstring(f'<p:transition xmlns:p="{P_NS}" spd="slow"><p:fade/></p:transition>'))
    s2.append(etree.fromstring(_timing([_click_group(_effect("emph", id2["EMPHASIS STAYS"]))])))
    parts["ppt/slides/slide2.xml"] = etree.tostring(s2, xml_declaration=True, encoding="UTF-8", standalone=True)

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in infos:
            zout.writestr(info, parts[info.filename])


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else OUT
    build_deck(out)
    inject(out)
    print(f"wrote {out}")
